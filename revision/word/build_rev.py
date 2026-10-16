"""Build the revised Word manuscript, with revised text highlighted.

Adapted from JCLLIT_SUBMISSION/word/build_src.py. It flattens
paper/jcllt_revision.tex into pandoc-friendly LaTeX (tables inlined, numeric
macros and cross-references resolved to the values the compiled PDF shows,
citations replaced by their numbers), turns every \\rev{...} into markers that
revmarks.lua converts into a highlighted character style, runs pandoc against
the journal's Word template, and gives that style a yellow highlight.

    cd revision/word && python3 build_rev.py

Run pdflatex on paper/jcllt_revision.tex first so the .aux file is current.
"""
import os, re, subprocess

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
PAPER = os.path.join(ROOT, "paper")

src = open(os.path.join(PAPER, "jcllt_revision.tex"), encoding="utf-8").read()
bib = open(os.path.join(PAPER, "bib_jcllt_rev.tex"), encoding="utf-8").read()
nums = {}
for f in ("numbers.tex", "numbers_rev.tex"):
    nums.update(dict(re.findall(r"\\newcommand\{\\(\w+)\}\{(.*)\}", open(os.path.join(PAPER, f)).read())))
labels = dict(re.findall(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", open(os.path.join(PAPER, "jcllt_revision.aux")).read()))
order = re.findall(r"\\bibitem\{([^}]*)\}", bib)
citenum = {k: str(i + 1) for i, k in enumerate(order)}


def braced(s, i):
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{" and (j == 0 or s[j - 1] != "\\"):
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced brace")


def unwrap(s, cmd, open_tok, close_tok):
    """Replace \\cmd{X} by open_tok X close_tok, innermost-safe, left to right."""
    key = "\\" + cmd + "{"
    while key in s:
        i = s.index(key)
        inner, j = braced(s, i + len(key) - 1)
        s = s[:i] + open_tok + inner + close_tok + s[j:]
    return s


src = re.sub(r"(?m)(?<!\\)%.*$", "", src)
body = src[src.index("\\section{Introduction}"):src.index("\\input{bib_jcllt_rev}")]
abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", src, re.S).group(1).strip()
keywords = re.search(r"\\begin\{keywords\}(.*?)\\end\{keywords\}", src, re.S).group(1).strip()

front = r"""\textbf{RESEARCH ARTICLE}

Journal of Computational Law and Legal Technology

\begin{center}
{\Large \textbf{\textsc{LexFlip}: A Dissociation Diagnostic for Legal Meaning Preservation Metrics}}
\end{center}

Gaurab Baral\textsuperscript{1,*}

\textsuperscript{1}\,Department of Computer Science, University of Cincinnati, United States, baralgb@mail.uc.edu.

\textbf{*Corresponding author:} Gaurab Baral, Department of Computer Science, University of Cincinnati, United States. Email: baralgb@mail.uc.edu

\textbf{Abstract:} ABSTRACTBODY

\textbf{Keywords:} KEYWORDBODY

""".replace("ABSTRACTBODY", abstract).replace("KEYWORDBODY", keywords)
body = front + body


def inline_table(m):
    t = open(os.path.join(PAPER, "tables", m.group(1) + ".tex"), encoding="utf-8").read()
    return re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}\s*", "", t)


body = re.sub(r"\\input\{tables/(\w+)\}", inline_table, body)

_sec = [0, 0]


def number_heading(m):
    level, title = m.group(1), m.group(2)
    if level == "section":
        _sec[0] += 1; _sec[1] = 0
        return "\\section{%d. %s}" % (_sec[0], title)
    _sec[1] += 1
    return "\\subsection{%d.%d. %s}" % (_sec[0], _sec[1], title)


body = re.sub(r"\\(section|subsection)\{([^}]*)\}(?!\*)", number_heading, body)
for k, v in sorted(nums.items(), key=lambda kv: -len(kv[0])):
    body = body.replace("\\" + k + "\\ ", v + " ")
    body = body.replace("\\" + k + "{}", v)
    body = re.sub(r"\\" + k + r"(?![a-zA-Z])", lambda m, v=v: v, body)
body = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels.get(m.group(1), "??"), body)
body = re.sub(r"\\cite\{([^}]*)\}",
              lambda m: "[" + ", ".join(citenum.get(k.strip(), "?") for k in m.group(1).split(",")) + "]", body)

FLOAT = re.compile(r"\\begin\{(table|figure)(\*?)\}(.*?)\\end\{\1\2\}", re.S)


def float_block(m):
    kind, inner = m.group(1), m.group(3)
    lab = re.search(r"\\label\{([^}]*)\}", inner)
    num = labels.get(lab.group(1), "?") if lab else "?"
    ci = inner.index("\\caption{") + len("\\caption")
    cap, after = braced(inner, ci)
    rest = inner[:inner.index("\\caption{")] + inner[after:]
    rest = re.sub(r"\\label\{[^}]*\}", "", rest).replace("\\centering", "")
    rest = re.sub(r"\\(small|footnotesize|scriptsize)\b", "", rest)
    while "\\resizebox" in rest:
        i = rest.index("\\resizebox") + len("\\resizebox")
        _, i = braced(rest, i); _, i = braced(rest, i)
        inner_box, j = braced(rest, i)
        rest = rest[:rest.index("\\resizebox")] + inner_box + rest[j:]
    rest = re.sub(r"\\includegraphics\[[^\]]*\]\{figs/(\w+)\.pdf\}", r"\\includegraphics{figs/\1.png}", rest)
    name = "Table" if kind == "table" else "Figure"
    cap = " ".join(cap.split())
    return "\n\n\\textbf{%s %s.} %s\n\n%s\n\n" % (name, num, cap, rest.strip())


body = FLOAT.sub(float_block, body)
body = body.replace("\\flushbottom", "").replace("\\balance", "")
body = re.sub(r"\\(begingroup|endgroup)\b", "", body)
body = re.sub(r"\\setlength\{[^}]*\}\{[^}]*\}", "", body)
body = re.sub(r"\\AtBeginEnvironment\{[^}]*\}\{[^}]*\}", "", body)
body = re.sub(r"\\label\{[^}]*\}", "", body)
body = re.sub(r"\\itemsep\s*\d+pt", "", body)
body = unwrap(body, "pending", r"\textbf{[PENDING: ", "]}")

# \rev{...} -> markers. A \rev that spans paragraphs or list environments is
# fine: the Lua filter keeps its state across blocks, in document order.
CLEAN = os.environ.get("CLEAN") == "1"   # plain journal format, no highlighting
body = unwrap(body, "rev", "", "") if CLEAN else unwrap(body, "rev", " REVSTARTX ", " REVENDX ")

entries = re.split(r"\\bibitem\{[^}]*\}", bib)[1:]
refs = ["\\section*{References}\n"]
for i, e in enumerate(entries):
    e = e.split("\\end{thebibliography}")[0]
    e = re.sub(r"^\s*%\s*", "", e)
    refs.append("[%d] %s\n" % (i + 1, " ".join(e.split())))
body += "\n\n" + "\n\n".join(refs) + "\n"
doc = ("\\documentclass{article}\n\\usepackage{graphicx}\n\\usepackage{booktabs}\n"
       "\\usepackage{amsmath,amssymb}\n\\begin{document}\n" + body + "\n\\end{document}\n")
doc = doc.replace("\\noindent ", "")
open(os.path.join(HERE, "manuscript_rev_src.tex"), "w", encoding="utf-8").write(doc)

out = os.path.join(ROOT, "revision", "JCLLT_LexFlip_manuscript_revised%s.docx" % ("_clean" if CLEAN else "_highlighted"))
subprocess.run(["pandoc", "manuscript_rev_src.tex", "-f", "latex", "-t", "docx",
                "--reference-doc=JCLLT Template-OTH.docx", "--lua-filter=revmarks.lua",
                "--resource-path=.:" + PAPER, "-o", out], cwd=HERE, check=True)

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
d = Document(out)
n = 0
for st in d.styles:
    if st.name == "Revised":
        st.font.highlight_color = WD_COLOR_INDEX.YELLOW
        n += 1
leftover = sum(p.text.count("REVSTARTX") + p.text.count("REVENDX") for p in d.paragraphs)
d.save(out)
print("wrote", out, "| Revised style found:", bool(n), "| leftover markers:", leftover,
      "| pending:", sum(p.text.count("[PENDING") for p in d.paragraphs))
