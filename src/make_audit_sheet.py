"""Build the expert audit workbook for LEXFLIP and its specificity controls.

Requested in review (all four reviewers): an audit of a stratified sample of
the definitional labels by Quebec jurists, with agreement statistics, and a
check of each sentence in the context of its article, since provisions depend
on cross-references and defined terms.

The workbook is blind: an annotator sees the article, the original sentence and
the edited sentence, never the rule, the tier or the cell the item came from.
Items are shuffled. The key that maps item ids back to cells and rules is
written to a separate CSV that is not sent to annotators.

Sample (stratified, seed 12050):
  LEXFLIP     6 per perturbation rule (all items where a rule has fewer)
  neutral     24, spread over the neutral rules
  paraphrase  24, if the paraphrase cell exists
"""
import os, sys, json, re, random, csv
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import probes as P
import specificity as SP
from augment import DATA

OUT = os.path.join(P.ROOT, "audit")
os.makedirs(OUT, exist_ok=True)
SEED = 12050
PER_RULE, N_NEUTRAL, N_PARA = 6, 24, 24


def articles():
    arts = []
    for f, law in (("A-25.jsonl", "Loi sur l'assurance automobile (RLRQ c A-25)"),
                   ("code_securite_routiere.jsonl", "Code de la sécurité routière (RLRQ c C-24.2)")):
        for line in open(os.path.join(DATA, f), encoding="utf-8"):
            t = json.loads(line)["text"]
            arts.append((law, t, re.sub(r"\s+", " ", t.replace("\n", " "))))
    return arts


ARTS = articles()


def context(sentence):
    for law, raw, flat in ARTS:
        if sentence in flat:
            m = re.match(r"\s*(\d+(?:\.\d+)*)\.", raw)
            return law, ("art. " + m.group(1)) if m else "", raw.strip()
    return "", "", ""


def stratified(items, key, per, rng):
    by = {}
    for it in items:
        by.setdefault(key(it), []).append(it)
    out = []
    for k in sorted(by):
        g = by[k][:]; rng.shuffle(g); out += g[:per]
    return out


def main():
    rng = random.Random(SEED)
    lex = [dict(cell="lexflip", rule=r, tier=int(P.T1[i] and 1 or 2), src=i,
                original=P.SOURCES[i], edited=P.edit[i]["simplification"])
           for i, r in enumerate(P.RULES)]
    sample = stratified(lex, lambda d: d["rule"], PER_RULE, rng)

    NE = SP.neutral_edits(P.SOURCES)
    neu = [dict(cell="neutral", rule=r, tier="", src=i, original=P.SOURCES[i], edited=e)
           for i, e, r in NE]
    per = max(1, N_NEUTRAL // len({d["rule"] for d in neu}))
    s_neu = stratified(neu, lambda d: d["rule"], per, rng)
    rest = [d for d in neu if d not in s_neu]; rng.shuffle(rest)
    sample += (s_neu + rest)[:N_NEUTRAL]

    try:
        import gen_paraphrases as G
        para = [dict(cell="paraphrase", rule="paraphrase", tier="", src=i,
                     original=P.SOURCES[i], edited=p) for i, p in G.load()]
        rng.shuffle(para); sample += para[:N_PARA]
    except FileNotFoundError:
        print("  no paraphrase cell yet; the sheet carries LEXFLIP and neutral items only")

    rng.shuffle(sample)
    for k, d in enumerate(sample, 1):
        d["item"] = "Q%03d" % k
        d["law"], d["article"], d["article_text"] = context(d["original"])

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ins = wb.active; ins.title = "Instructions"
    text = [
        "Audit d'experts / Expert audit: LexFlip",
        "",
        "FR. Chaque ligne de la feuille « Items » présente une phrase tirée d'une loi québécoise "
        "(colonne « Phrase originale »), l'article complet dont elle provient, et une version modifiée "
        "de cette phrase. Pour chaque ligne, veuillez répondre aux quatre questions en vous plaçant "
        "dans le contexte de l'article et de la loi (définitions, renvois). Travaillez seul(e), sans "
        "consulter les autres annotateurs. Il n'y a pas de bonne réponse attendue : certaines "
        "modifications changent la portée juridique, d'autres non.",
        "",
        "EN. Each row of the 'Items' sheet shows a sentence from a Quebec statute, the full article it "
        "comes from, and an edited version of the sentence. For each row, answer the four questions, "
        "reading the sentence in the context of its article and statute (definitions, cross-references). "
        "Work independently. There is no expected answer: some edits change the legal effect and some do not.",
        "",
        "Q1  Qualité de la langue / Language: is the edited sentence acceptable statutory French? "
        "(1 = unacceptable ... 5 = fully natural)",
        "Q2  La modification change-t-elle la portée juridique ? / Does the edit change the legal effect? "
        "(Oui / Non / Incertain)",
        "Q3  Si oui, de quelle façon ? / If yes, how? (Inverse la règle / Modifie la portée sans l'inverser / Autre)",
        "Q4  Gravité / Severity of the change for a person relying on the edited text "
        "(0 = none ... 4 = decisive)",
        "Commentaire / Comment: optional, e.g. a defined term or cross-reference that matters.",
    ]
    for r, t in enumerate(text, 1):
        c = ins.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1:
            c.font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Items")
    head = ["Item", "Loi / Statute", "Article", "Article complet / Full article",
            "Phrase originale / Original sentence", "Phrase modifiée / Edited sentence",
            "Q1 Langue (1-5)", "Q2 Portée changée ?", "Q3 Nature", "Q4 Gravité (0-4)",
            "Commentaire / Comment"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="FADF70")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for d in sample:
        ws.append([d["item"], d["law"], d["article"], d["article_text"], d["original"],
                   d["edited"], None, None, None, None, None])
    widths = [7, 22, 9, 60, 45, 45, 10, 12, 18, 10, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(sample) + 1
    for col, opts in (("G", '"1,2,3,4,5"'), ("H", '"Oui,Non,Incertain"'),
                      ("I", '"Inverse la règle,Modifie la portée sans l\'inverser,Autre"'),
                      ("J", '"0,1,2,3,4"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True)
        ws.add_data_validation(dv); dv.add("%s2:%s%d" % (col, col, n))
    ws.freeze_panes = "E2"
    path = os.path.join(OUT, "LexFlip_expert_audit.xlsx")
    wb.save(path)

    with open(os.path.join(OUT, "audit_key.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["item", "cell", "rule", "tier", "source_index", "article_found"])
        for d in sample:
            w.writerow([d["item"], d["cell"], d["rule"], d["tier"], d["src"], bool(d["article"])])
    import collections
    print("wrote %s (%d items): %s" % (path, len(sample), dict(collections.Counter(d["cell"] for d in sample))))
    print("article context found for %d / %d" % (sum(bool(d["article"]) for d in sample), len(sample)))


if __name__ == "__main__":
    main()
